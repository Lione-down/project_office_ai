from copy import deepcopy
from hashlib import sha256
from pathlib import Path
import shutil
from uuid import UUID

from openpyxl import load_workbook
import pytest

from app.services.quality.spreadsheet.contracts import SpreadsheetQualityConflict
from app.services.quality.spreadsheet.inspector import SpreadsheetQualityInspector
from app.services.renderers.excel.contracts import WorkbookPlan
from app.services.renderers.excel.libreoffice import (
    FormulaCacheDifference,
    WorkbookValidationReport,
)
from tests.renderers.spreadsheet.test_contracts import valid_workbook_plan
from tests.renderers.spreadsheet.test_renderer import render_sample


ASSET_ID = UUID("019f6000-0000-7000-8000-000000000011")


def validation_report(
    rendered,
    validation_copy: Path,
    *,
    differences: tuple[FormulaCacheDifference, ...] = (),
) -> WorkbookValidationReport:
    return WorkbookValidationReport(
        schema_version="workbook-validation-report.v1",
        original_sha256=rendered.sha256,
        validation_sha256=sha256(validation_copy.read_bytes()).hexdigest(),
        plan_sha256=rendered.plan_sha256,
        renderer_version=rendered.renderer_version,
        template_code=rendered.template_code,
        template_version=rendered.template_version,
        libreoffice_version="LibreOffice 25.2.3",
        formula_count=12,
        difference_count=len(differences),
        differences=differences,
    )


def inspect_copy(
    original: Path,
    validation_copy: Path,
    *,
    differences: tuple[FormulaCacheDifference, ...] = (),
    conflicts: tuple[SpreadsheetQualityConflict, ...] = (),
):
    rendered = render_sample(original)
    plan = WorkbookPlan.validate_for_renderer(valid_workbook_plan())
    report = validation_report(rendered, validation_copy, differences=differences)
    return SpreadsheetQualityInspector().inspect(
        asset_id=ASSET_ID,
        rendered=rendered,
        plan=plan,
        validation_copy=validation_copy,
        validation_report=report,
        validation_report_sha256="e" * 64,
        conflicts=conflicts,
    )


@pytest.mark.parametrize(
    ("error_value", "expected_code"),
    [
        ("#REF!", "formula_error_ref"),
        ("#VALUE!", "formula_error_value"),
        ("#DIV/0!", "formula_error_div_zero"),
        ("#NAME?", "formula_error_name"),
    ],
)
def test_formula_error_values_block_download(
    tmp_path: Path, error_value: str, expected_code: str
) -> None:
    original = tmp_path / "original.xlsx"
    render_sample(original)
    validation_copy = tmp_path / "validation.xlsx"
    shutil.copyfile(original, validation_copy)
    workbook = load_workbook(validation_copy)
    workbook["经营看板"]["B3"] = error_value
    workbook.save(validation_copy)
    workbook.close()

    report = inspect_copy(
        tmp_path / f"source-{expected_code}.xlsx",
        validation_copy,
    )

    assert report.status == "blocked"
    assert report.download_allowed is False
    assert expected_code in {finding.code for finding in report.findings}


def test_formula_error_like_plain_text_is_not_blocked(tmp_path: Path) -> None:
    original = tmp_path / "original.xlsx"
    render_sample(original)
    validation_copy = tmp_path / "validation.xlsx"
    shutil.copyfile(original, validation_copy)
    workbook = load_workbook(validation_copy)
    cell = workbook["销售明细"]["A3"]
    cell.value = "#REF!"
    cell.data_type = "s"
    workbook.save(validation_copy)
    workbook.close()

    report = inspect_copy(tmp_path / "source.xlsx", validation_copy)

    assert "formula_error_ref" not in {finding.code for finding in report.findings}


def test_raw_numeric_value_drift_from_plan_is_a_field_conflict(tmp_path: Path) -> None:
    original = tmp_path / "original.xlsx"
    render_sample(original)
    validation_copy = tmp_path / "validation.xlsx"
    shutil.copyfile(original, validation_copy)
    workbook = load_workbook(validation_copy)
    workbook["销售明细"]["D3"] = 999
    workbook.save(validation_copy)
    workbook.close()

    report = inspect_copy(tmp_path / "source.xlsx", validation_copy)

    assert report.status == "blocked"
    assert "field_value_conflict" in {finding.code for finding in report.findings}


def test_workbook_that_cannot_open_is_blocked(tmp_path: Path) -> None:
    original = tmp_path / "original.xlsx"
    rendered = render_sample(original)
    plan = WorkbookPlan.validate_for_renderer(valid_workbook_plan())
    validation_copy = tmp_path / "validation.xlsx"
    validation_copy.write_bytes(b"not-an-xlsx")
    report = validation_report(rendered, validation_copy)

    inspected = SpreadsheetQualityInspector().inspect(
        asset_id=ASSET_ID,
        rendered=rendered,
        plan=plan,
        validation_copy=validation_copy,
        validation_report=report,
        validation_report_sha256="e" * 64,
    )

    assert inspected.status == "blocked"
    assert "workbook_open_failed" in {finding.code for finding in inspected.findings}


def test_deterministic_recalculation_mismatch_blocks_without_formula_text(
    tmp_path: Path,
) -> None:
    original = tmp_path / "original.xlsx"
    render_sample(original)
    validation_copy = tmp_path / "validation.xlsx"
    shutil.copyfile(original, validation_copy)
    differences = (
        FormulaCacheDifference(
            sheet_name="经营看板",
            cell_reference="B3",
            formula_changed=False,
            original_cached_value=300,
            recalculated_cached_value=301,
        ),
    )

    report = inspect_copy(
        tmp_path / "source.xlsx",
        validation_copy,
        differences=differences,
    )

    assert report.status == "blocked"
    assert "deterministic_recalculation_mismatch" in {
        finding.code for finding in report.findings
    }
    assert "=SUM(" not in report.model_dump_json()


def test_chart_range_outside_existing_cells_blocks_download(tmp_path: Path) -> None:
    original = tmp_path / "original.xlsx"
    render_sample(original)
    validation_copy = tmp_path / "validation.xlsx"
    shutil.copyfile(original, validation_copy)
    workbook = load_workbook(validation_copy)
    workbook["经营看板"]._charts[0].ser[0].val.numRef.f = "销售明细!$Z$3:$Z$5"
    workbook.save(validation_copy)
    workbook.close()

    report = inspect_copy(tmp_path / "source.xlsx", validation_copy)

    assert report.status == "blocked"
    assert "chart_range_invalid" in {finding.code for finding in report.findings}


def test_chart_range_pointing_to_wrong_existing_field_blocks_download(
    tmp_path: Path,
) -> None:
    original = tmp_path / "original.xlsx"
    render_sample(original)
    validation_copy = tmp_path / "validation.xlsx"
    shutil.copyfile(original, validation_copy)
    workbook = load_workbook(validation_copy)
    workbook["经营看板"]._charts[0].ser[0].val.numRef.f = "销售明细!$C$3:$C$5"
    workbook.save(validation_copy)
    workbook.close()

    report = inspect_copy(tmp_path / "source.xlsx", validation_copy)

    assert report.status == "blocked"
    assert "chart_range_invalid" in {finding.code for finding in report.findings}


def test_validation_copy_size_limit_fails_closed_before_workbook_open(
    tmp_path: Path,
) -> None:
    original = tmp_path / "original.xlsx"
    rendered = render_sample(original)
    plan = WorkbookPlan.validate_for_renderer(valid_workbook_plan())
    validation_copy = tmp_path / "validation.xlsx"
    validation_copy.write_bytes(b"x" * 17)
    report = validation_report(rendered, validation_copy)

    inspected = SpreadsheetQualityInspector(max_workbook_bytes=16).inspect(
        asset_id=ASSET_ID,
        rendered=rendered,
        plan=plan,
        validation_copy=validation_copy,
        validation_report=report,
        validation_report_sha256="e" * 64,
    )

    assert inspected.status == "blocked"
    assert "workbook_file_too_large" in {finding.code for finding in inspected.findings}


@pytest.mark.parametrize(
    ("mutation", "conflicts", "expected_code"),
    [
        (
            lambda workbook: setattr(workbook["销售明细"], "title", "未知明细"),
            (),
            "sheet_contract_conflict",
        ),
        (
            lambda workbook: setattr(workbook["销售明细"]["D3"], "value", "not-number"),
            (),
            "field_type_conflict",
        ),
        (
            lambda _workbook: None,
            (
                SpreadsheetQualityConflict(
                    kind="unit",
                    code="unit_conflict",
                    resolved=False,
                ),
            ),
            "unit_conflict",
        ),
    ],
)
def test_sheet_field_and_unit_conflicts_are_blockers(
    tmp_path: Path,
    mutation,
    conflicts: tuple[SpreadsheetQualityConflict, ...],
    expected_code: str,
) -> None:
    original = tmp_path / "original.xlsx"
    render_sample(original)
    validation_copy = tmp_path / "validation.xlsx"
    shutil.copyfile(original, validation_copy)
    workbook = load_workbook(validation_copy)
    mutation(workbook)
    workbook.save(validation_copy)
    workbook.close()

    report = inspect_copy(
        tmp_path / f"source-{expected_code}.xlsx",
        validation_copy,
        conflicts=conflicts,
    )

    assert report.status == "blocked"
    assert expected_code in {finding.code for finding in report.findings}


def test_warning_only_report_remains_downloadable(tmp_path: Path) -> None:
    original = tmp_path / "original.xlsx"
    rendered = render_sample(original)
    validation_copy = tmp_path / "validation.xlsx"
    shutil.copyfile(original, validation_copy)
    plan_payload = deepcopy(valid_workbook_plan())
    plan = WorkbookPlan.validate_for_renderer(plan_payload)
    report = validation_report(rendered, validation_copy)

    inspected = SpreadsheetQualityInspector().inspect(
        asset_id=ASSET_ID,
        rendered=rendered,
        plan=plan,
        validation_copy=validation_copy,
        validation_report=report,
        validation_report_sha256="e" * 64,
        conflicts=(
            SpreadsheetQualityConflict(
                kind="field",
                code="field_label_warning",
                resolved=True,
            ),
        ),
    )

    assert inspected.status == "passed_with_warnings"
    assert inspected.download_allowed is True
