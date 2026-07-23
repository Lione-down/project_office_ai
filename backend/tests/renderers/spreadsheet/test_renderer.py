from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
import pytest

from app.services.renderers.excel.contracts import WorkbookPlan
from app.services.renderers.excel.formulas import FORMULA_FUNCTION_WHITELIST
from app.services.renderers.excel.renderer import ExcelRenderer, ExcelRendererError
from tests.renderers.spreadsheet.test_contracts import valid_workbook_plan


def render_sample(path: Path, template_code: str = "excel.analysis-dashboard.v1"):
    plan = WorkbookPlan.validate_for_renderer(valid_workbook_plan(template_code))
    return ExcelRenderer().render(plan, path)


@pytest.mark.parametrize(
    "template_code",
    [
        "excel.analysis-dashboard.v1",
        "excel.operating-review.v1",
    ],
)
def test_two_system_templates_generate_real_traceable_xlsx(
    tmp_path: Path, template_code: str
) -> None:
    path = tmp_path / f"{template_code}.xlsx"

    rendered = render_sample(path, template_code)

    assert rendered.path == path.resolve()
    assert rendered.byte_size == path.stat().st_size
    assert rendered.sha256 == sha256(path.read_bytes()).hexdigest()
    assert rendered.plan_sha256
    assert rendered.schema_version == "workbook-plan.v1"
    assert rendered.renderer_version == "excel-renderer.v1"
    assert rendered.template_code == template_code
    assert rendered.template_version == "1.0.0"
    assert rendered.sheet_count == 2
    assert rendered.generated_cell_count == 38
    assert ZipFile(path).testzip() is None


def test_renderer_writes_expected_english_formulas_and_cached_values(
    tmp_path: Path,
) -> None:
    path = tmp_path / "analysis.xlsx"
    render_sample(path)

    formulas = load_workbook(path, data_only=False, read_only=False, keep_links=False)
    cached = load_workbook(path, data_only=True, read_only=True, keep_links=False)
    data_formulas = formulas["销售明细"]
    data_cached = cached["销售明细"]
    summary_formulas = formulas["经营看板"]
    summary_cached = cached["经营看板"]

    assert data_formulas["E3"].value == "=IFERROR((D3-C3)/C3,0)"
    assert data_formulas["F4"].value == "=IFERROR((D4-C4)/C4,0)"
    assert data_formulas["G3"].value == "=RANK.EQ(D3,$D$3:$D$5,0)"
    assert data_cached["E3"].value == pytest.approx(0.2)
    assert data_cached["F4"].value == pytest.approx(-0.25)
    assert data_cached["G3"].value == 1
    assert data_cached["G4"].value == 3
    assert data_cached["G5"].value == 1

    assert summary_formulas["B3"].value == "=SUM('销售明细'!$D$3:$D$5)"
    assert summary_formulas["B4"].value == "=AVERAGE('销售明细'!$D$3:$D$5)"
    assert summary_formulas["B5"].value == "=COUNT('销售明细'!$D$3:$D$5)"
    assert summary_cached["B3"].value == 300
    assert summary_cached["B4"].value == 100
    assert summary_cached["B5"].value == 3
    assert FORMULA_FUNCTION_WHITELIST == frozenset(
        {"AVERAGE", "COUNT", "IFERROR", "RANK.EQ", "SUM"}
    )

    formulas.close()
    cached.close()


def test_renderer_writes_chart_format_validation_freeze_and_print_contract(
    tmp_path: Path,
) -> None:
    path = tmp_path / "analysis.xlsx"
    render_sample(path)

    workbook = load_workbook(path, data_only=False, read_only=False, keep_links=False)
    data = workbook["销售明细"]
    summary = workbook["经营看板"]
    chart = summary._charts[0]

    assert data.freeze_panes == "A3"
    assert summary.freeze_panes == "A3"
    assert len(data.conditional_formatting) == 1
    assert data.data_validations.count == 1
    assert data.print_options.horizontalCentered is True
    assert data.page_setup.orientation == "landscape"
    assert data.sheet_properties.pageSetUpPr.fitToPage is True
    assert data.page_setup.fitToWidth in {None, 1}
    assert data.page_setup.fitToHeight == 0
    assert data.print_title_rows == "$2:$2"
    assert len(summary._charts) == 1
    assert chart.ser[0].val.numRef.f == "销售明细!$D$3:$D$5"
    assert chart.ser[0].cat.numRef.f == "销售明细!$A$3:$A$5"

    workbook.close()


def test_renderer_treats_untrusted_formula_and_url_like_strings_as_text(
    tmp_path: Path,
) -> None:
    path = tmp_path / "analysis.xlsx"
    render_sample(path)

    workbook = load_workbook(path, data_only=False, read_only=False, keep_links=False)
    cell = workbook["销售明细"]["A5"]

    assert cell.value == "=UNTRUSTED()"
    assert cell.data_type == "s"
    with ZipFile(path) as archive:
        names = set(archive.namelist())
        assert not any(name.startswith("xl/externalLinks/") for name in names)
        assert "xl/vbaProject.bin" not in names

    workbook.close()


def test_renderer_is_byte_deterministic_for_same_plan(tmp_path: Path) -> None:
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"

    first_result = render_sample(first)
    second_result = render_sample(second)

    assert first.read_bytes() == second.read_bytes()
    assert first_result.sha256 == second_result.sha256
    assert first_result.plan_sha256 == second_result.plan_sha256


def test_renderer_never_overwrites_existing_destination(tmp_path: Path) -> None:
    destination = tmp_path / "existing.xlsx"
    destination.write_bytes(b"existing-user-data")

    with pytest.raises(
        ExcelRendererError,
        match="workbook_output_already_exists",
    ):
        render_sample(destination)

    assert destination.read_bytes() == b"existing-user-data"
