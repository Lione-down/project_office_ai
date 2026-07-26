from __future__ import annotations

from hashlib import sha256
import json
import math
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from app.services.golden.contracts import (
    Artifact,
    ExpectedMetric,
    GoldenCaseManifest,
    GoldenRegressionError,
    GoldenRegressionReport,
)
from app.services.golden.manifest import load_golden_manifest
from app.services.golden.security import inspect_office_artifact
from app.services.quality.presentation.inspector import PresentationPackageInspector
from app.services.renderers.excel.contracts import WorkbookPlan
from app.services.renderers.excel.renderer import ExcelRenderer


def workbook_structural_sha256(path: str | Path) -> str:
    digest = sha256()
    try:
        with ZipFile(path) as package:
            entries = package.infolist()
            names = [entry.filename for entry in entries]
            if len(names) != len(set(names)):
                raise GoldenRegressionError("golden_spreadsheet_package_invalid")
            for name in sorted(names):
                content = package.read(name)
                if name.endswith((".xml", ".rels")):
                    content = content.replace(b"\r\n", b"\n")
                encoded_name = name.encode("utf-8")
                digest.update(len(encoded_name).to_bytes(4, "big"))
                digest.update(encoded_name)
                digest.update(len(content).to_bytes(8, "big"))
                digest.update(content)
    except (BadZipFile, OSError) as exception:
        raise GoldenRegressionError("golden_spreadsheet_package_invalid") from exception
    return digest.hexdigest()


class GoldenRegressionRunner:
    def __init__(self, *, repo_root: str | Path) -> None:
        self._repo_root = Path(repo_root).resolve()

    def verify_artifact(
        self,
        path: str | Path,
        *,
        expected_sha256: str,
        expected_byte_size: int,
    ) -> None:
        artifact = Path(path)
        if not artifact.is_file() or artifact.stat().st_size != expected_byte_size:
            raise GoldenRegressionError("golden_artifact_integrity_failed")
        digest = sha256(artifact.read_bytes()).hexdigest()
        if digest != expected_sha256:
            raise GoldenRegressionError("golden_artifact_integrity_failed")

    def run(self, manifest_path: str | Path) -> GoldenRegressionReport:
        manifest = load_golden_manifest(manifest_path, repo_root=self._repo_root)
        for artifact in manifest.tracked_files:
            self._verify_declared_artifact(artifact)
        for artifact in manifest.artifacts:
            inspection = inspect_office_artifact(self._repo_root / artifact.path)
            if not inspection.is_safe:
                raise GoldenRegressionError("golden_artifact_security_failed")
        self._verify_requirements(manifest)
        self._verify_metrics(
            self._repo_root / manifest.input_workbook.path,
            manifest.expected_metrics,
        )
        workbook_plan = WorkbookPlan.validate_for_renderer(
            json.loads(
                (self._repo_root / manifest.workbook_plan.path).read_text(
                    encoding="utf-8"
                )
            )
        )
        spreadsheet_runs = self._repeat_spreadsheet_generation(
            workbook_plan,
            manifest.repetitions.spreadsheet,
            manifest.expected_workbook,
        )
        presentation_runs = self._repeat_presentation_reference(
            self._repo_root / manifest.expected_presentation.path,
            repetitions=manifest.repetitions.presentation,
            expected_slide_count=manifest.expected_presentation.slide_count or 0,
            evidence_ids=set(manifest.expected_evidence_ids),
        )
        return GoldenRegressionReport(
            schema_version="golden-regression-report.v1",
            case_id=manifest.case_id,
            coverage_label=manifest.coverage.label,
            status="passed",
            spreadsheet_generation_runs=spreadsheet_runs,
            presentation_reference_runs=presentation_runs,
            presentation_generation_runs=0,
            presentation_generation_evidence="ppt_renderer_suite_required",
            findings=(),
        )

    def _verify_declared_artifact(self, artifact: Artifact) -> None:
        self.verify_artifact(
            self._repo_root / artifact.path,
            expected_sha256=artifact.sha256,
            expected_byte_size=artifact.byte_size,
        )

    def _verify_requirements(self, manifest: GoldenCaseManifest) -> None:
        page_count = manifest.expected_presentation.slide_count
        sheet_count = manifest.expected_workbook.sheet_count
        page_minimum, page_maximum = manifest.requirements.presentation_page_range
        sheet_minimum, sheet_maximum = manifest.requirements.workbook_sheet_range
        if (
            page_count is None
            or not page_minimum <= page_count <= page_maximum
            or sheet_count is None
            or not sheet_minimum <= sheet_count <= sheet_maximum
        ):
            raise GoldenRegressionError("golden_requirement_mismatch")
        workbook = load_workbook(
            self._repo_root / manifest.expected_workbook.path,
            data_only=False,
            read_only=False,
            keep_links=False,
        )
        try:
            chart_count = sum(len(sheet._charts) for sheet in workbook.worksheets)
        finally:
            workbook.close()
        if chart_count != manifest.requirements.workbook_chart_count:
            raise GoldenRegressionError("golden_requirement_mismatch")

    def _verify_metrics(
        self,
        path: Path,
        metrics: tuple[ExpectedMetric, ...],
    ) -> None:
        workbook = load_workbook(path, data_only=True, read_only=True, keep_links=False)
        try:
            for metric in metrics:
                if metric.source_sheet_index > len(workbook.worksheets):
                    raise GoldenRegressionError("golden_metric_mismatch")
                worksheet = workbook.worksheets[metric.source_sheet_index - 1]
                min_col, min_row, max_col, max_row = range_boundaries(
                    metric.source_range
                )
                columns = [
                    [
                        worksheet.cell(row=row, column=column).value
                        for row in range(min_row, max_row + 1)
                    ]
                    for column in range(min_col, max_col + 1)
                ]
                values = [
                    [self._number(value) for value in column] for column in columns
                ]
                actual = self._aggregate(metric.aggregation, values)
                if not math.isclose(
                    actual,
                    metric.value,
                    rel_tol=0,
                    abs_tol=metric.tolerance,
                ):
                    raise GoldenRegressionError("golden_metric_mismatch")
        finally:
            workbook.close()

    @staticmethod
    def _number(value) -> float:
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            raise GoldenRegressionError("golden_metric_mismatch")
        return float(value)

    @staticmethod
    def _aggregate(kind: str, columns: list[list[float]]) -> float:
        if kind == "sum" and len(columns) == 1:
            return sum(columns[0])
        if kind == "max" and len(columns) == 1:
            return max(columns[0])
        if len(columns) != 2:
            raise GoldenRegressionError("golden_metric_mismatch")
        first = sum(columns[0])
        second = sum(columns[1])
        if kind == "ratio_of_sums" and first:
            return second / first
        if kind == "growth_of_sums" and second:
            return (first - second) / second
        raise GoldenRegressionError("golden_metric_mismatch")

    def _repeat_spreadsheet_generation(
        self,
        plan: WorkbookPlan,
        repetitions: int,
        expected: Artifact,
    ) -> int:
        renderer = ExcelRenderer()
        expected_structural_sha256 = workbook_structural_sha256(
            self._repo_root / expected.path
        )
        with TemporaryDirectory(prefix="office-ai-golden-xlsx-") as directory:
            for index in range(repetitions):
                target = Path(directory) / f"run-{index + 1}.xlsx"
                rendered = renderer.render(plan, target)
                if (
                    workbook_structural_sha256(rendered.path)
                    != expected_structural_sha256
                    or rendered.sheet_count != expected.sheet_count
                ):
                    raise GoldenRegressionError("golden_spreadsheet_drift")
        return repetitions

    @staticmethod
    def _repeat_presentation_reference(
        path: Path,
        *,
        repetitions: int,
        expected_slide_count: int,
        evidence_ids: set[str],
    ) -> int:
        inspector = PresentationPackageInspector()
        for _ in range(repetitions):
            result = inspector.inspect(
                path,
                expected_slide_count=expected_slide_count,
                expected_evidence_ids=evidence_ids,
            )
            if result.findings:
                raise GoldenRegressionError("golden_presentation_reference_failed")
        return repetitions
