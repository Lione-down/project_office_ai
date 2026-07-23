from dataclasses import dataclass
from hashlib import sha256
import json
import os
from pathlib import Path
import shutil
import subprocess
from typing import Annotated, Literal

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field, ValidationError, model_validator

from app.services.renderers.excel.renderer import RenderedWorkbook


CachedValue = str | int | float | bool | None


class StrictModel(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)


class FormulaCacheDifference(StrictModel):
    sheet_name: Annotated[str, Field(min_length=1, max_length=31)]
    cell_reference: Annotated[str, Field(pattern=r"^[A-Z]{1,3}[1-9][0-9]{0,6}$")]
    formula_changed: bool
    original_cached_value: CachedValue
    recalculated_cached_value: CachedValue


class WorkbookValidationReport(StrictModel):
    schema_version: Literal["workbook-validation-report.v1"]
    original_sha256: Annotated[str, Field(pattern=r"^[0-9a-f]{64}$")]
    validation_sha256: Annotated[str, Field(pattern=r"^[0-9a-f]{64}$")]
    plan_sha256: Annotated[str, Field(pattern=r"^[0-9a-f]{64}$")]
    renderer_version: Literal["excel-renderer.v1"]
    template_code: Literal[
        "excel.analysis-dashboard.v1",
        "excel.operating-review.v1",
    ]
    template_version: Literal["1.0.0"]
    libreoffice_version: Annotated[str, Field(min_length=1, max_length=64)]
    formula_count: Annotated[int, Field(ge=0, le=200_000)]
    difference_count: Annotated[int, Field(ge=0, le=200_000)]
    differences: Annotated[
        tuple[FormulaCacheDifference, ...],
        Field(max_length=200_000),
    ]

    @model_validator(mode="after")
    def counts_match_payload(self):
        if self.difference_count != len(self.differences):
            raise ValueError("difference count does not match payload")
        return self


class WorkbookValidationError(RuntimeError):
    def __init__(self, error_code: str) -> None:
        super().__init__(error_code)
        self.error_code = error_code


@dataclass(frozen=True)
class WorkbookValidationArtifacts:
    validation_copy_path: Path
    report_path: Path
    report: WorkbookValidationReport


@dataclass(frozen=True)
class _FormulaSnapshot:
    formula_sha256: str
    cached_value: CachedValue


class LibreOfficeWorkbookValidator:
    def __init__(
        self,
        root: str | Path,
        *,
        soffice_executable: str = "soffice",
        runner=subprocess.run,
        timeout_seconds: float = 120,
        max_bytes: int = 200 * 1024 * 1024,
    ) -> None:
        self._root = Path(root).resolve()
        self._soffice = soffice_executable
        self._runner = runner
        self._timeout = timeout_seconds
        self._max_bytes = max_bytes

    def validate(self, rendered: RenderedWorkbook) -> WorkbookValidationArtifacts:
        original = rendered.path.resolve()
        original_digest, original_size = self._hash_file(original)
        if original_digest != rendered.sha256 or original_size != rendered.byte_size:
            raise WorkbookValidationError("workbook_original_integrity_failed")
        original_formulas = self._formula_snapshots(original)

        self._root.mkdir(parents=True, exist_ok=True)
        validation_directory = self._root / rendered.sha256
        validation_output = validation_directory / "recalculated" / "analysis.xlsx"
        report_path = validation_directory / "validation-report.json"
        if validation_directory.exists():
            return self._replay(
                rendered,
                original_formulas,
                validation_output,
                report_path,
            )
        try:
            validation_directory.mkdir(exist_ok=False)
        except FileExistsError:
            return self._replay(
                rendered,
                original_formulas,
                validation_output,
                report_path,
            )
        input_directory = validation_directory / "input"
        output_directory = validation_directory / "recalculated"
        input_directory.mkdir()
        output_directory.mkdir()
        validation_input = input_directory / "analysis.xlsx"
        try:
            shutil.copyfile(original, validation_input)
            version = self._libreoffice_version()
            self._recalculate(validation_input, output_directory)
            if not validation_output.is_file():
                raise WorkbookValidationError("libreoffice_output_missing")
            validation_digest, _ = self._hash_file(validation_output)
            current_digest, current_size = self._hash_file(original)
            if current_digest != rendered.sha256 or current_size != rendered.byte_size:
                raise WorkbookValidationError("workbook_original_integrity_failed")
            validation_formulas = self._formula_snapshots(validation_output)
            differences = self._compare_formula_snapshots(
                original_formulas,
                validation_formulas,
            )
            report = WorkbookValidationReport(
                schema_version="workbook-validation-report.v1",
                original_sha256=rendered.sha256,
                validation_sha256=validation_digest,
                plan_sha256=rendered.plan_sha256,
                renderer_version=rendered.renderer_version,
                template_code=rendered.template_code,
                template_version=rendered.template_version,
                libreoffice_version=version,
                formula_count=len(original_formulas),
                difference_count=len(differences),
                differences=tuple(differences),
            )
            final_digest, final_size = self._hash_file(original)
            if final_digest != rendered.sha256 or final_size != rendered.byte_size:
                raise WorkbookValidationError("workbook_original_integrity_failed")
            self._write_report(report_path, report)
            return WorkbookValidationArtifacts(
                validation_copy_path=validation_output,
                report_path=report_path,
                report=report,
            )
        except Exception:
            shutil.rmtree(validation_directory, ignore_errors=True)
            raise

    def _replay(
        self,
        rendered: RenderedWorkbook,
        original_formulas: dict[tuple[str, str], _FormulaSnapshot],
        validation_output: Path,
        report_path: Path,
    ) -> WorkbookValidationArtifacts:
        if not validation_output.is_file() or not report_path.is_file():
            raise WorkbookValidationError("workbook_validation_incomplete")
        try:
            if report_path.stat().st_size > 10 * 1024 * 1024:
                raise WorkbookValidationError("workbook_validation_integrity_failed")
            report = WorkbookValidationReport.model_validate_json(
                report_path.read_bytes()
            )
        except WorkbookValidationError:
            raise
        except (OSError, ValidationError, ValueError) as exception:
            raise WorkbookValidationError(
                "workbook_validation_integrity_failed"
            ) from exception
        validation_digest, _ = self._hash_file(validation_output)
        if (
            report.original_sha256 != rendered.sha256
            or report.validation_sha256 != validation_digest
            or report.plan_sha256 != rendered.plan_sha256
            or report.renderer_version != rendered.renderer_version
            or report.template_code != rendered.template_code
            or report.template_version != rendered.template_version
        ):
            raise WorkbookValidationError("workbook_validation_integrity_failed")
        recalculated = self._formula_snapshots(validation_output)
        differences = tuple(
            self._compare_formula_snapshots(original_formulas, recalculated)
        )
        if (
            report.formula_count != len(original_formulas)
            or report.difference_count != len(differences)
            or report.differences != differences
        ):
            raise WorkbookValidationError("workbook_validation_integrity_failed")
        return WorkbookValidationArtifacts(
            validation_copy_path=validation_output,
            report_path=report_path,
            report=report,
        )

    def _libreoffice_version(self) -> str:
        try:
            completed = self._runner(
                [self._soffice, "--version"],
                check=True,
                capture_output=True,
                text=True,
                timeout=self._timeout,
            )
        except FileNotFoundError as exception:
            raise WorkbookValidationError("libreoffice_unavailable") from exception
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as exception:
            raise WorkbookValidationError("libreoffice_unavailable") from exception
        version = str(completed.stdout).strip()
        if not version:
            raise WorkbookValidationError("libreoffice_version_invalid")
        return version[:64]

    def _recalculate(self, source: Path, output_directory: Path) -> None:
        profile_directory = output_directory.parent / "profile"
        try:
            self._runner(
                [
                    self._soffice,
                    "--headless",
                    "--nologo",
                    "--nodefault",
                    "--nofirststartwizard",
                    "--nolockcheck",
                    "--norestore",
                    f"-env:UserInstallation={profile_directory.as_uri()}",
                    "--convert-to",
                    "xlsx",
                    "--outdir",
                    str(output_directory),
                    str(source),
                ],
                check=True,
                capture_output=True,
                text=True,
                timeout=self._timeout,
            )
        except FileNotFoundError as exception:
            raise WorkbookValidationError("libreoffice_unavailable") from exception
        except subprocess.TimeoutExpired as exception:
            raise WorkbookValidationError("libreoffice_timeout") from exception
        except subprocess.CalledProcessError as exception:
            raise WorkbookValidationError(
                "libreoffice_recalculation_failed"
            ) from exception

    def _hash_file(self, path: Path) -> tuple[str, int]:
        if not path.is_file():
            raise WorkbookValidationError("workbook_file_missing")
        digest = sha256()
        byte_size = 0
        with path.open("rb") as handle:
            while chunk := handle.read(1024 * 1024):
                byte_size += len(chunk)
                if byte_size > self._max_bytes:
                    raise WorkbookValidationError("workbook_file_too_large")
                digest.update(chunk)
        return digest.hexdigest(), byte_size

    @staticmethod
    def _formula_snapshots(path: Path) -> dict[tuple[str, str], _FormulaSnapshot]:
        try:
            formulas = load_workbook(
                path,
                data_only=False,
                read_only=True,
                keep_links=False,
            )
            cached = load_workbook(
                path,
                data_only=True,
                read_only=True,
                keep_links=False,
            )
        except Exception as exception:
            raise WorkbookValidationError(
                "workbook_formula_snapshot_invalid"
            ) from exception
        snapshots: dict[tuple[str, str], _FormulaSnapshot] = {}
        try:
            if formulas.sheetnames != cached.sheetnames:
                raise WorkbookValidationError("workbook_formula_snapshot_invalid")
            for formula_sheet, cached_sheet in zip(
                formulas.worksheets,
                cached.worksheets,
                strict=True,
            ):
                for formula_row, cached_row in zip(
                    formula_sheet.iter_rows(),
                    cached_sheet.iter_rows(),
                    strict=True,
                ):
                    for formula_cell, cached_cell in zip(
                        formula_row,
                        cached_row,
                        strict=True,
                    ):
                        if formula_cell.data_type != "f":
                            continue
                        formula = str(formula_cell.value)
                        snapshots[(formula_sheet.title, formula_cell.coordinate)] = (
                            _FormulaSnapshot(
                                formula_sha256=sha256(
                                    formula.encode("utf-8")
                                ).hexdigest(),
                                cached_value=cached_cell.value,
                            )
                        )
                        if len(snapshots) > 200_000:
                            raise WorkbookValidationError(
                                "workbook_formula_limit_exceeded"
                            )
        except WorkbookValidationError:
            raise
        except Exception as exception:
            raise WorkbookValidationError(
                "workbook_formula_snapshot_invalid"
            ) from exception
        finally:
            formulas.close()
            cached.close()
        return snapshots

    @staticmethod
    def _compare_formula_snapshots(
        original: dict[tuple[str, str], _FormulaSnapshot],
        recalculated: dict[tuple[str, str], _FormulaSnapshot],
    ) -> list[FormulaCacheDifference]:
        differences = []
        for key in sorted(set(original) | set(recalculated)):
            before = original.get(key)
            after = recalculated.get(key)
            if before == after:
                continue
            differences.append(
                FormulaCacheDifference(
                    sheet_name=key[0],
                    cell_reference=key[1],
                    formula_changed=(
                        before is None
                        or after is None
                        or before.formula_sha256 != after.formula_sha256
                    ),
                    original_cached_value=(
                        None if before is None else before.cached_value
                    ),
                    recalculated_cached_value=(
                        None if after is None else after.cached_value
                    ),
                )
            )
        return differences

    @staticmethod
    def _write_report(
        report_path: Path,
        report: WorkbookValidationReport,
    ) -> None:
        payload = json.dumps(
            report.model_dump(mode="json"),
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
        with report_path.open("xb") as handle:
            handle.write(payload)
            handle.flush()
            os.fsync(handle.fileno())
