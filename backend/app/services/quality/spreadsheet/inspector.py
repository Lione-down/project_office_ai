from datetime import UTC, datetime
from hashlib import sha256
from pathlib import Path
import re
from uuid import UUID

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string

from app.services.quality.spreadsheet.contracts import (
    SpreadsheetQualityConflict,
    SpreadsheetQualityFinding,
    SpreadsheetQualityReport,
    SpreadsheetSourceReference,
    SpreadsheetValidationReference,
)
from app.services.renderers.excel.contracts import (
    ComparisonCalculation,
    DataSheet,
    RankCalculation,
    SummarySheet,
    WorkbookPlan,
    workbook_plan_sha256,
)
from app.services.renderers.excel.formulas import (
    comparison_formula,
    rank_formula,
    summary_formula,
)
from app.services.renderers.excel.libreoffice import WorkbookValidationReport
from app.services.renderers.excel.renderer import RenderedWorkbook


ERROR_CODES = {
    "#REF!": "formula_error_ref",
    "#VALUE!": "formula_error_value",
    "#DIV/0!": "formula_error_div_zero",
    "#NAME?": "formula_error_name",
}
RANGE_PATTERN = re.compile(
    r"^'?([^']+(?:''[^']*)*)'?!\$([A-Z]{1,3})\$(\d+):\$([A-Z]{1,3})\$(\d+)$"
)


class SpreadsheetQualityInspector:
    def __init__(self, *, max_workbook_bytes: int = 200 * 1024 * 1024) -> None:
        self._max_workbook_bytes = max_workbook_bytes

    def inspect(
        self,
        *,
        asset_id: UUID,
        rendered: RenderedWorkbook,
        plan: WorkbookPlan,
        validation_copy: str | Path,
        validation_report: WorkbookValidationReport,
        validation_report_sha256: str,
        conflicts: tuple[SpreadsheetQualityConflict, ...] = (),
    ) -> SpreadsheetQualityReport:
        findings: dict[str, SpreadsheetQualityFinding] = {}

        def add(code: str, severity: str, rule_version: str) -> None:
            findings.setdefault(
                code,
                SpreadsheetQualityFinding(
                    code=code,
                    severity=severity,
                    rule_version=rule_version,
                ),
            )

        validation_path = Path(validation_copy).resolve()
        validation_size = validation_path.stat().st_size
        can_open = validation_size <= self._max_workbook_bytes
        if can_open:
            validation_sha256 = self._digest(validation_path)
        else:
            validation_sha256 = validation_report.validation_sha256
            add(
                "workbook_file_too_large",
                "blocker",
                "spreadsheet-size.v1",
            )
        if not self._trace_matches(
            rendered, plan, validation_report, validation_sha256
        ):
            add(
                "workbook_validation_integrity_failed",
                "blocker",
                "spreadsheet-validation-integrity.v1",
            )
        for conflict in conflicts:
            add(
                conflict.code,
                "warning" if conflict.resolved else "blocker",
                f"spreadsheet-{conflict.kind}-conflict.v1",
            )
        for difference in validation_report.differences:
            add(
                "deterministic_recalculation_mismatch",
                "blocker",
                "spreadsheet-recalculation.v1",
            )
            if difference.formula_changed:
                add(
                    "formula_definition_changed",
                    "blocker",
                    "spreadsheet-formula-integrity.v1",
                )
        if can_open:
            formulas = cached = None
            try:
                formulas = load_workbook(
                    validation_path,
                    data_only=False,
                    read_only=False,
                    keep_links=False,
                )
                cached = load_workbook(
                    validation_path,
                    data_only=True,
                    read_only=False,
                    keep_links=False,
                )
                self._inspect_workbook(plan, formulas, cached, add)
            except Exception:
                add(
                    "workbook_open_failed",
                    "blocker",
                    "spreadsheet-open.v1",
                )
            finally:
                if formulas is not None:
                    formulas.close()
                if cached is not None:
                    cached.close()
        ordered = tuple(findings[key] for key in sorted(findings))
        has_blocker = any(item.severity == "blocker" for item in ordered)
        has_warning = any(item.severity == "warning" for item in ordered)
        status = (
            "blocked"
            if has_blocker
            else "passed_with_warnings"
            if has_warning
            else "passed"
        )
        return SpreadsheetQualityReport(
            schema_version="workbook-quality-report.v1",
            quality_engine_version="spreadsheet-quality.v1",
            source=SpreadsheetSourceReference(
                asset_id=asset_id,
                sha256=rendered.sha256,
                byte_size=rendered.byte_size,
                plan_sha256=rendered.plan_sha256,
                renderer_version=rendered.renderer_version,
                template_code=rendered.template_code,
                template_version=rendered.template_version,
            ),
            validation=SpreadsheetValidationReference(
                status="passed",
                sha256=validation_sha256,
                byte_size=validation_size,
                report_sha256=validation_report_sha256,
                libreoffice_version=validation_report.libreoffice_version,
                formula_count=validation_report.formula_count,
                difference_count=validation_report.difference_count,
            ),
            status=status,
            download_allowed=not has_blocker,
            findings=ordered,
            generated_at=datetime.now(UTC),
        )

    def _inspect_workbook(self, plan, formulas, cached, add) -> None:
        expected_names = [sheet.name for sheet in plan.sheets]
        if formulas.sheetnames != expected_names or cached.sheetnames != expected_names:
            add(
                "sheet_contract_conflict",
                "blocker",
                "spreadsheet-sheet-contract.v1",
            )
            return
        self._scan_formula_errors(cached, add)
        for sheet in plan.sheets:
            if isinstance(sheet, DataSheet):
                self._inspect_data_sheet(sheet, cached[sheet.name], add)
            else:
                self._inspect_summary_sheet(
                    sheet,
                    plan,
                    formulas[sheet.name],
                    cached,
                    add,
                )

    @staticmethod
    def _scan_formula_errors(workbook, add) -> None:
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    code = (
                        ERROR_CODES.get(cell.value) if cell.data_type == "e" else None
                    )
                    if code is not None:
                        add(code, "blocker", "spreadsheet-formula-errors.v1")

    def _inspect_data_sheet(self, plan: DataSheet, worksheet, add) -> None:
        columns = [*plan.columns, *plan.computed_columns]
        if [worksheet.cell(2, index + 1).value for index in range(len(columns))] != [
            item.header for item in columns
        ]:
            add(
                "field_contract_conflict",
                "blocker",
                "spreadsheet-field-contract.v1",
            )
        for row_index, row in enumerate(plan.rows, start=3):
            for column_index, column in enumerate(plan.columns, start=1):
                actual = worksheet.cell(row_index, column_index).value
                if not self._value_matches(actual, column.value_type):
                    add(
                        "field_type_conflict",
                        "blocker",
                        "spreadsheet-field-types.v1",
                    )
                elif not self._values_equal(actual, row[column.key]):
                    add(
                        "field_value_conflict",
                        "blocker",
                        "spreadsheet-field-values.v1",
                    )
        self._check_computed_values(plan, worksheet, add)

    def _check_computed_values(self, plan: DataSheet, worksheet, add) -> None:
        indexes = {
            column.key: index
            for index, column in enumerate([*plan.columns, *plan.computed_columns])
        }
        first_row = 3
        last_row = first_row + len(plan.rows) - 1
        for computed in plan.computed_columns:
            target = indexes[computed.key] + 1
            calculation = computed.calculation
            if isinstance(calculation, ComparisonCalculation):
                for offset, row in enumerate(plan.rows):
                    _, expected = comparison_formula(
                        calculation,
                        first_row + offset,
                        indexes,
                        row,
                    )
                    self._compare_value(
                        worksheet.cell(first_row + offset, target).value,
                        expected,
                        add,
                    )
            elif isinstance(calculation, RankCalculation):
                values = tuple(row[calculation.source_column] for row in plan.rows)
                for offset in range(len(plan.rows)):
                    _, expected = rank_formula(
                        calculation,
                        first_row + offset,
                        first_row,
                        last_row,
                        indexes,
                        values,
                        offset,
                    )
                    self._compare_value(
                        worksheet.cell(first_row + offset, target).value,
                        expected,
                        add,
                    )

    def _inspect_summary_sheet(
        self,
        plan: SummarySheet,
        workbook_plan: WorkbookPlan,
        formula_sheet,
        cached_workbook,
        add,
    ) -> None:
        data_sheets = {
            sheet.name: sheet
            for sheet in workbook_plan.sheets
            if isinstance(sheet, DataSheet)
        }
        for offset, metric in enumerate(plan.metrics, start=3):
            source = data_sheets[metric.calculation.source_sheet]
            indexes = {
                column.key: index
                for index, column in enumerate(
                    [*source.columns, *source.computed_columns]
                )
            }
            values = tuple(row[metric.calculation.source_column] for row in source.rows)
            _, expected = summary_formula(
                metric.calculation,
                indexes[metric.calculation.source_column],
                3,
                2 + len(source.rows),
                values,
            )
            self._compare_value(
                cached_workbook[plan.name].cell(offset, 2).value,
                expected,
                add,
            )
        if len(formula_sheet._charts) != len(plan.charts):
            add(
                "chart_range_invalid",
                "blocker",
                "spreadsheet-chart-range.v1",
            )
            return
        for chart, chart_plan in zip(formula_sheet._charts, plan.charts, strict=True):
            for series in chart.ser:
                references = []
                if series.val is not None and series.val.numRef is not None:
                    references.append(series.val.numRef.f)
                if series.cat is not None:
                    if series.cat.numRef is not None:
                        references.append(series.cat.numRef.f)
                    if series.cat.strRef is not None:
                        references.append(series.cat.strRef.f)
                if any(
                    not self._range_exists(cached_workbook, reference)
                    for reference in references
                ):
                    add(
                        "chart_range_invalid",
                        "blocker",
                        "spreadsheet-chart-range.v1",
                    )
            if self._chart_ranges(chart) != self._expected_chart_ranges(
                workbook_plan, chart_plan
            ):
                add(
                    "chart_range_invalid",
                    "blocker",
                    "spreadsheet-chart-range.v1",
                )

    @staticmethod
    def _range_exists(workbook, reference: str) -> bool:
        parsed = SpreadsheetQualityInspector._parse_range(reference)
        if parsed is None:
            return False
        sheet_name, first_column, first_row, last_column, last_row = parsed
        if sheet_name not in workbook.sheetnames:
            return False
        worksheet = workbook[sheet_name]
        return (
            first_column <= last_column <= worksheet.max_column
            and first_row <= last_row <= worksheet.max_row
        )

    @staticmethod
    def _chart_ranges(chart) -> set[tuple[str, int, int, int, int]]:
        references: list[str] = []
        for series in chart.ser:
            if series.val is not None and series.val.numRef is not None:
                references.append(series.val.numRef.f)
            if series.cat is not None:
                if series.cat.numRef is not None:
                    references.append(series.cat.numRef.f)
                if series.cat.strRef is not None:
                    references.append(series.cat.strRef.f)
        return {
            parsed
            for reference in references
            if (parsed := SpreadsheetQualityInspector._parse_range(reference))
            is not None
        }

    @staticmethod
    def _expected_chart_ranges(
        workbook_plan: WorkbookPlan, chart_plan
    ) -> set[tuple[str, int, int, int, int]]:
        source = next(
            sheet
            for sheet in workbook_plan.sheets
            if isinstance(sheet, DataSheet) and sheet.name == chart_plan.source_sheet
        )
        columns = {
            column.key: index
            for index, column in enumerate(
                [*source.columns, *source.computed_columns], start=1
            )
        }
        first_row = 3
        last_row = 2 + len(source.rows)
        return {
            (
                source.name,
                columns[column],
                first_row,
                columns[column],
                last_row,
            )
            for column in [
                chart_plan.categories_column,
                *chart_plan.value_columns,
            ]
        }

    @staticmethod
    def _parse_range(
        reference: str,
    ) -> tuple[str, int, int, int, int] | None:
        matched = RANGE_PATTERN.fullmatch(reference)
        if matched is None:
            return None
        return (
            matched.group(1).replace("''", "'"),
            column_index_from_string(matched.group(2)),
            int(matched.group(3)),
            column_index_from_string(matched.group(4)),
            int(matched.group(5)),
        )

    @staticmethod
    def _value_matches(value, value_type: str) -> bool:
        if value is None:
            return True
        if value_type == "text":
            return isinstance(value, str)
        if value_type == "number":
            return isinstance(value, (int, float)) and not isinstance(value, bool)
        if value_type == "integer":
            return isinstance(value, int) and not isinstance(value, bool)
        return isinstance(value, bool)

    @staticmethod
    def _compare_value(actual, expected, add) -> None:
        if not SpreadsheetQualityInspector._values_equal(actual, expected):
            add(
                "deterministic_recalculation_mismatch",
                "blocker",
                "spreadsheet-recalculation.v1",
            )

    @staticmethod
    def _values_equal(actual, expected) -> bool:
        if isinstance(actual, (int, float)) and isinstance(expected, (int, float)):
            return abs(float(actual) - float(expected)) <= max(
                1e-9, abs(float(expected)) * 1e-9
            )
        return actual == expected

    @staticmethod
    def _trace_matches(
        rendered: RenderedWorkbook,
        plan: WorkbookPlan,
        report: WorkbookValidationReport,
        validation_sha256: str,
    ) -> bool:
        return (
            report.original_sha256 == rendered.sha256
            and report.validation_sha256 == validation_sha256
            and report.plan_sha256 == rendered.plan_sha256
            and report.renderer_version == rendered.renderer_version
            and report.template_code == rendered.template_code
            and report.template_version == rendered.template_version
            and rendered.plan_sha256 == workbook_plan_sha256(plan)
        )

    def _digest(self, path: Path) -> str:
        digest = sha256()
        size = 0
        with path.open("rb") as handle:
            while chunk := handle.read(1024 * 1024):
                size += len(chunk)
                if size > self._max_workbook_bytes:
                    raise ValueError("workbook_file_too_large")
                digest.update(chunk)
        return digest.hexdigest()
